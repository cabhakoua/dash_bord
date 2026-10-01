"""
Import des fichiers déposés depuis l'ordinateur.

  - offres : classeur mensuel .xlsx (même structure que offres.xlsx) ou .csv
             (colonnes date, offre, activations ; ou date + une colonne par offre)
  - churn  : .csv de même structure que churn_database2.csv

Les règles reprennent celles des scripts 02, 03 et 05 du projet.
"""
import datetime as dt
import io
import re
import unicodedata

import numpy as np
import pandas as pd

# ============================================================== OFFRES
MOIS = {
    "Offers April": (2022, 4), "Offers May": (2022, 5), "Offers June": (2022, 6),
    "Offers July": (2022, 7), "Offers August": (2022, 8), "September": (2022, 9),
    "October": (2022, 10), "November": (2022, 11), "December": (2022, 12),
    "Jan23": (2023, 1), "Feb23": (2023, 2), "Mar23": (2023, 3), "Apr23": (2023, 4),
    "May23": (2023, 5), "Jun23": (2023, 6), "Jul23": (2023, 7), "Aug23": (2023, 8),
    "Sept23": (2023, 9), "Oct23": (2023, 10), "Nov23": (2023, 11), "Dec23": (2023, 12),
    "Janv24": (2024, 1), "Feb24": (2024, 2), "Mar24": (2024, 3), "Avr24": (2024, 4),
    "Mai24": (2024, 5), "Jun24": (2024, 6), "Jul24": (2024, 7), "Aou24": (2024, 8),
    "Sept24": (2024, 9), "Oct24": (2024, 10), "Nov24": (2024, 11), "Dec24": (2024, 12),
    "JanV25": (2025, 1), "Feb25": (2025, 2), "Mars25": (2025, 3), "Avr25": (2025, 4),
    "Mai25": (2025, 5), "Juin25": (2025, 6), "Juillet25": (2025, 7), "Août25": (2025, 8),
}
ALIAS = {
    "spot": "Blue mo S", "spot / blue mo s": "Blue mo S",
    "cool": "Blue mo M", "cool / blue mo m": "Blue mo M",
    "blue booster": "Blue mo L", "blue booster / blue mo l": "Blue mo L",
    "kolo": "Blue mo XL", "kolo / blue mo xl": "Blue mo XL",
    "swim": "Blue mo XXL", "swim / blue mo xxl": "Blue mo XXL",
    "fap": "FAP",
}
FORFAITS_DEFAUT = ["Blue One S", "Blue One M", "Blue One L"]


def _norm_label(s):
    if s is None:
        return None
    s = unicodedata.normalize("NFKD", str(s).strip()).encode("ascii", "ignore").decode()
    s = re.sub(r"\s+", " ", s).strip()
    if s.lower() in ALIAS:
        return ALIAS[s.lower()]
    if " / " in s:
        return s.split(" / ")[-1].strip()
    return s


def _to_num(v):
    if v is None:
        return np.nan
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    s = str(v).strip()
    if s in {"", "-", "--"} or s.startswith("="):
        return np.nan
    s = s.replace("\u202f", "").replace(" ", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return np.nan


def _cle_libelle(s):
    """Clé de comparaison : minuscules, sans espaces ni ponctuation."""
    return re.sub(r"[^a-z0-9]", "", str(s).lower())


def unifier_libelles(df):
    """
    Fusionne les offres écrites de plusieurs façons (« FAKO Week » / « FAKO WEEK »,
    « X-tremNet+15 Go_CAMTEL » / « X-tremNet+15Go_CAMTEL »). On garde l'écriture
    qui pèse le plus en activations. Retourne (table, journal).
    """
    cles = df["offre"].map(_cle_libelle)
    poids = df.groupby(["offre"])["activations"].sum()
    choix, journal = {}, []
    for cle, groupe in pd.Series(df["offre"].unique()).groupby(
            pd.Series(df["offre"].unique()).map(_cle_libelle)):
        variantes = list(groupe)
        if len(variantes) > 1:
            retenu = max(variantes, key=lambda v: poids.get(v, 0))
            for v in variantes:
                choix[v] = retenu
            journal.append("Écritures fusionnées : " + " + ".join(sorted(variantes))
                           + f" -> {retenu}")
    if choix:
        df = df.copy()
        df["offre"] = df["offre"].replace(choix)
        df = df.groupby(["date", "offre"], as_index=False)["activations"].sum()
    return df, journal


def familles(offres):
    """
    Regroupe les offres par gamme (« Blue One S », « Blue One M » -> « Blue One »).
    Retourne un dictionnaire offre -> gamme. Sert aux raccourcis de sélection.
    """
    vus, res = {}, {}
    for o in offres:
        bas = str(o).lower().strip()
        mots = re.split(r"[\s_]+", str(o).strip())
        if "mboa" in bas:
            g = "MBOA"
        elif bas.startswith("fako"):
            g = "FAKO"
        elif bas.startswith("toli"):
            g = "Toli"
        elif bas.startswith(("x-", "special x", "pack xtrem", "deal", "enjoy", "skate")):
            g = "X-Trem et offres spéciales"
        elif bas.startswith("camtel_") or bas.startswith(("cb pro", "ptt")):
            g = "Autres"
        elif bas.startswith("blue") and len(mots) > 1:
            cle = " ".join(mots[:2]).lower()
            g = vus.setdefault(cle, " ".join(mots[:2]))
        else:
            g = "Autres"
        res[o] = g
    return res


def resume_offres(long):
    """Une ligne par offre : début, fin, jours renseignés, total, couverture."""
    g = long.groupby("offre").agg(debut=("date", "min"), fin=("date", "max"),
                                  jours=("date", "nunique"),
                                  total=("activations", "sum"))
    g["couverture"] = (g["jours"] / ((g["fin"] - g["debut"]).dt.days + 1)).round(2)
    return g.sort_values("total", ascending=False)


def lire_classeur(data: bytes):
    """Classeur mensuel -> table longue (date, offre, activations) + journal."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    mois_connus = {k.strip(): v for k, v in MOIS.items()}
    records, journal = [], []

    for sheet in wb.sheetnames:
        cle = sheet.strip()
        if cle not in mois_connus:
            journal.append(f"Feuille ignorée (mois non reconnu) : {sheet!r}")
            continue
        annee, mois = mois_connus[cle]
        rows = list(wb[sheet].iter_rows(values_only=True))
        if not rows:
            continue
        header = rows[0]
        col_lab = 0
        for j, v in enumerate(header[:3]):
            if isinstance(v, str) and "number of activations" in v.lower():
                col_lab = j
                break

        date_cols = {}
        for j in range(col_lab + 1, len(header)):
            v = header[j]
            if isinstance(v, (dt.datetime, dt.date)):
                d = v.date() if isinstance(v, dt.datetime) else v
                if (d.year, d.month) != (annee, mois):
                    try:
                        d = dt.date(annee, mois, d.day)
                    except ValueError:
                        continue
                date_cols[j] = d
            elif isinstance(v, str) and v.strip().upper() == "TOTAL":
                break
        if date_cols:
            j0, d0 = min(date_cols.items())
            for j in range(col_lab + 1, len(header)):
                v = header[j]
                if j in date_cols:
                    continue
                if isinstance(v, str) and v.strip().startswith("="):
                    cand = d0 + dt.timedelta(days=j - j0)
                    if cand.month == mois:
                        date_cols[j] = cand
                elif isinstance(v, str) and v.strip().upper() == "TOTAL":
                    break

        for r in rows[1:]:
            if r is None or len(r) <= col_lab or r[col_lab] is None:
                continue
            lab = str(r[col_lab]).strip()
            low = lab.lower()
            if lab.upper() == "TOTAL" or "sim card" in low or "number of activations" in low:
                continue
            offre = _norm_label(lab)
            if not offre:
                continue
            for j, d in date_cols.items():
                if j < len(r):
                    val = _to_num(r[j])
                    if not np.isnan(val):
                        records.append((d, offre, val))

    if not records:
        raise ValueError("Aucune activation lisible dans ce classeur. Vérifiez que les "
                         "noms de feuilles suivent la convention du fichier d'origine.")
    df = pd.DataFrame(records, columns=["date", "offre", "activations"])
    df["date"] = pd.to_datetime(df["date"])
    df = df.groupby(["date", "offre"], as_index=False)["activations"].sum()
    df, fusion = unifier_libelles(df)
    journal.extend(fusion)
    journal.append(f"{df['offre'].nunique()} offres, {df['date'].nunique()} jours, "
                   f"du {df['date'].min().date()} au {df['date'].max().date()}")
    return df.sort_values(["offre", "date"]).reset_index(drop=True), journal


def lire_csv_offres(data: bytes):
    """CSV long (date, offre, activations) ou large (date + une colonne par offre)."""
    df = pd.read_csv(io.BytesIO(data), sep=None, engine="python")
    df.columns = [str(c).strip() for c in df.columns]
    bas = {c.lower(): c for c in df.columns}
    col_date = bas.get("date") or df.columns[0]
    df[col_date] = pd.to_datetime(df[col_date], dayfirst=True, errors="coerce")
    df = df.dropna(subset=[col_date])
    if "offre" in bas and "activations" in bas:
        long = df.rename(columns={col_date: "date", bas["offre"]: "offre",
                                  bas["activations"]: "activations"})
        long = long[["date", "offre", "activations"]]
    else:
        long = df.melt(id_vars=col_date, var_name="offre", value_name="activations")
        long = long.rename(columns={col_date: "date"})
    long["activations"] = pd.to_numeric(long["activations"], errors="coerce")
    long = long.dropna(subset=["activations"])
    long["offre"] = long["offre"].map(_norm_label)
    long = long.groupby(["date", "offre"], as_index=False)["activations"].sum()
    if long.empty:
        raise ValueError("Aucune ligne exploitable dans ce CSV.")
    long, fusion = unifier_libelles(long)
    return long, fusion + [f"{long['offre'].nunique()} offres, {long['date'].nunique()} jours"]


def lire_offres(nom: str, data: bytes):
    if nom.lower().endswith((".xlsx", ".xlsm")):
        return lire_classeur(data)
    return lire_csv_offres(data)


def series_propres(long, offres, debut, fin=None, fenetre=15, k=3.0, couverture_min=0.85):
    """
    Série quotidienne par offre, sur la fenêtre [debut, fin].

    - Une offre bien renseignée (au moins `couverture_min` des jours) : jours manquants
      comblés par interpolation.
    - Une offre peu renseignée : l'absence de ligne signifie « aucune activation ce
      jour », les jours manquants valent donc 0.
    - Valeurs aberrantes corrigées par un filtre de Hampel (médiane glissante,
      seuil k x 1,4826 x MAD).
    - Avant son lancement ou après son arrêt, une offre vaut 0 : les offres de dates
      différentes peuvent ainsi être additionnées.
    Retourne (table corrigée, journal).
    """
    journal = []
    debut = pd.Timestamp(debut)
    fin_max = long["date"].max()
    fin = min(pd.Timestamp(fin), fin_max) if fin is not None else fin_max
    idx = pd.date_range(debut, fin, freq="D")
    cols = {}
    for o in offres:
        brut = long[long["offre"] == o].set_index("date")["activations"].astype(float)
        brut = brut[~brut.index.duplicated()].sort_index().loc[debut:fin]
        if brut.empty:
            cols[o] = pd.Series(0.0, index=idx)
            journal.append(f"{o} : aucune donnée sur la période choisie, comptée à 0")
            continue
        s = brut.asfreq("D")
        cov = float(s.notna().mean())
        n_manq = int(s.isna().sum())
        if cov < couverture_min:
            s = s.fillna(0.0)
            note_trous = f"{n_manq} jour(s) sans ligne comptés à 0 (offre à ventes espacées)"
        else:
            note_trous = f"{n_manq} jour(s) manquant(s) comblé(s)"
        med = s.rolling(fenetre, center=True, min_periods=5).median()
        mad = (s - med).abs().rolling(fenetre, center=True, min_periods=5).median()
        seuil = k * 1.4826 * mad
        aber = ((s - med).abs() > seuil.replace(0, np.nan)).fillna(False)
        c = s.copy()
        c[aber] = med[aber]
        c = c.interpolate("linear", limit_direction="both").round().clip(lower=0)
        c = c.reindex(idx)
        avant = int(c.isna().sum())
        c = c.fillna(0.0)
        cols[o] = c
        msg = f"{o} : {note_trous}, {int(aber.sum())} valeur(s) aberrante(s) corrigée(s)"
        if s.index.min() > debut + pd.Timedelta(days=7) or s.index.max() < fin - pd.Timedelta(days=7):
            msg += (f" ; données du {s.index.min():%d/%m/%Y} au {s.index.max():%d/%m/%Y}, "
                    f"{avant} jour(s) hors de cette plage comptés à 0")
        journal.append(msg)
    tab = pd.concat(cols, axis=1)
    tab.index.name = "date"
    journal.append(f"Table finale : {len(tab)} jours, {tab.index.min().date()} "
                   f"au {tab.index.max().date()}")
    return tab, journal


def plage_active(serie, fin_fenetre, tolerance=30, min_jours=200, min_non_nuls=30):
    """
    Décide si une série d'offre se prête à une prévision individuelle.
    Retourne (série rognée de ses zéros de tête, motif d'exclusion ou None).
    """
    non_nuls = serie[serie > 0]
    if non_nuls.empty:
        return serie.iloc[0:0], "aucune activation sur la période"
    s = serie.loc[non_nuls.index[0]:]
    if non_nuls.index[-1] < pd.Timestamp(fin_fenetre) - pd.Timedelta(days=tolerance):
        return s, f"offre arrêtée (dernière activation le {non_nuls.index[-1]:%d/%m/%Y})"
    if len(s) < min_jours:
        return s, f"historique trop court ({len(s)} jours)"
    if len(non_nuls) < min_non_nuls:
        return s, f"trop peu de jours d'activation ({len(non_nuls)})"
    return s, None


# ============================================================== CHURN
COLONNES_CHURN = ["CUSTOMER_ID", "PURGE_TIME_ATSGSN", "SUB_CREATE_TIME", "CHURN_STATUS",
                  "PLAN_NAME", "MONTHLY_CHARGE", "RECHARGE_FREQUENCY", "TENURE_MONTHS",
                  "AGE", "REGION", "GENDER"]


def verifier_churn(df):
    return [c for c in COLONNES_CHURN if c not in df.columns]


def preparer_churn(df):
    """Reprend le script 03 : fuite, censure, aberrants, variables dérivées."""
    log, n0 = [], len(df)
    df = df.copy()
    df["purge_dt"] = pd.to_datetime(df["PURGE_TIME_ATSGSN"], format="%d/%m/%Y %H:%M")
    df["crea_dt"] = pd.to_datetime(df["SUB_CREATE_TIME"], format="%d/%m/%Y %H:%M")
    ref = df["purge_dt"].max()
    log.append(f"Date de référence (dernière purge observée) : {ref.date()}")

    protegees = {"CUSTOMER_ID", "CHURN_STATUS", "purge_dt", "crea_dt"}
    const = [c for c in df.columns if df[c].nunique(dropna=False) <= 1 and c not in protegees]
    quasi = [c for c in ["TGPPAMBRMAXUL", "TGPPAMBRMAXDL"]
             if c in df.columns and df[c].value_counts(normalize=True).iloc[0] > 0.995]
    fuite = [c for c in ["PURGE_TIME_ATSGSN", "MONTHS_SINCE_PURGE", "PURGEDONMME"]
             if c in df.columns]
    redond = [c for c in ["TOTAL_CHARGE", "Latitude", "Longitude", "MSISDN"]
              if c in df.columns]
    df = df.drop(columns=const + quasi + fuite + redond)
    log.append(f"{len(const) + len(quasi)} colonne(s) sans information, "
               f"{len(fuite)} variable(s) de fuite, {len(redond)} redondance(s) retirées")

    avant = len(df)
    ok = (df["AGE"].between(15, 100) & (df["TENURE_MONTHS"] >= 0)
          & (df["RECHARGE_FREQUENCY"] >= 0) & (df["purge_dt"] >= df["crea_dt"]))
    df = df[ok].copy()
    log.append(f"{avant - len(df)} ligne(s) écartée(s) pour incohérence métier")
    for c in ["AGE", "TENURE_MONTHS", "RECHARGE_FREQUENCY"]:
        lo, hi = df[c].quantile([0.01, 0.99])
        df[c] = df[c].clip(lo, hi)

    seuil = ref - pd.Timedelta(days=60)
    avant = len(df)
    df = df[df["crea_dt"] <= seuil].copy()
    log.append(f"{avant - len(df)} abonné(s) créé(s) après le {seuil.date()} écarté(s) "
               "(censure à droite)")

    df["anciennete_jours"] = (ref - df["crea_dt"]).dt.days
    df["intensite_recharge"] = df["RECHARGE_FREQUENCY"] / df["TENURE_MONTHS"].clip(lower=1)
    df["revenu_cumule"] = df["MONTHLY_CHARGE"] * df["RECHARGE_FREQUENCY"]
    df["jamais_recharge"] = (df["RECHARGE_FREQUENCY"] == 0).astype(int)
    df["mois_creation"] = df["crea_dt"].dt.month
    df["tranche_age"] = pd.cut(df["AGE"], [0, 20, 25, 30, 120],
                               labels=["<=20", "21-25", "26-30", ">30"]).astype(str)
    df = df.drop(columns=["SUB_CREATE_TIME", "purge_dt", "crea_dt"])
    log.append(f"Jeu final : {len(df)} lignes ({len(df) / n0:.1%} de l'original), "
               f"taux de churn {df['CHURN_STATUS'].mean():.2%}")
    return df, log


def _lire_dates(s):
    d = pd.to_datetime(s, format="%d/%m/%Y %H:%M", errors="coerce")
    if d.isna().mean() > 0.5:                      # autre format : on retente jour d'abord
        d = pd.to_datetime(s, dayfirst=True, errors="coerce")
    return d


def departs_churn(df):
    """
    Un départ par ligne : (CUSTOMER_ID, PLAN_NAME, REGION, date_depart).

    La date de départ est celle de la dernière trace de l'abonné sur le réseau
    (PURGE_TIME_ATSGSN). Seuls les abonnés étiquetés partis (CHURN_STATUS = 1) sont
    gardés, après les mêmes filtres que `preparer_churn` : la somme des départs
    reste ainsi cohérente avec les abonnés scorés dans le reste de l'application.
    """
    df = df.copy()
    df["purge_dt"] = _lire_dates(df["PURGE_TIME_ATSGSN"])
    df["crea_dt"] = _lire_dates(df["SUB_CREATE_TIME"])
    ref = df["purge_dt"].max()
    ok = (df["AGE"].between(15, 100) & (df["TENURE_MONTHS"] >= 0)
          & (df["RECHARGE_FREQUENCY"] >= 0) & (df["purge_dt"] >= df["crea_dt"]))
    df = df[ok]
    df = df[df["crea_dt"] <= ref - pd.Timedelta(days=60)]
    df = df[df["CHURN_STATUS"] == 1]
    out = pd.DataFrame({
        "CUSTOMER_ID": df["CUSTOMER_ID"].values,
        "PLAN_NAME": df["PLAN_NAME"].values,
        "REGION": df["REGION"].fillna("Non renseignée").astype(str).str.strip().values,
        "date_depart": df["purge_dt"].dt.normalize().values})
    return out.reset_index(drop=True)


def entrainer_churn(df):
    """XGBoost sur 75 % des abonnés, AUC sur les 25 % restants, puis scores de tous."""
    from sklearn.compose import ColumnTransformer
    from sklearn.metrics import roc_auc_score
    from sklearn.model_selection import train_test_split
    from sklearn.pipeline import Pipeline
    from sklearn.preprocessing import OneHotEncoder
    from xgboost import XGBClassifier

    y, ids = df["CHURN_STATUS"].values, df["CUSTOMER_ID"].values
    X = df.drop(columns=["CHURN_STATUS", "CUSTOMER_ID"])
    num = X.select_dtypes(include=[np.number]).columns.tolist()
    cat = X.select_dtypes(exclude=[np.number]).columns.tolist()

    def pipe():
        return Pipeline([
            ("pre", ColumnTransformer([
                ("cat", OneHotEncoder(handle_unknown="ignore", drop="first"), cat),
                ("num", "passthrough", num)])),
            ("clf", XGBClassifier(n_estimators=400, max_depth=5, learning_rate=0.05,
                                  subsample=0.8, colsample_bytree=0.8,
                                  eval_metric="logloss", n_jobs=-1, random_state=42))])

    Xtr, Xte, ytr, yte = train_test_split(X, y, test_size=0.25, stratify=y, random_state=42)
    auc = float(roc_auc_score(yte, pipe().fit(Xtr, ytr).predict_proba(Xte)[:, 1]))
    final = pipe().fit(X, y)

    scores = pd.DataFrame({
        "CUSTOMER_ID": ids, "PLAN_NAME": df["PLAN_NAME"].values,
        "REGION": df["REGION"].values, "MONTHLY_CHARGE": df["MONTHLY_CHARGE"].values,
        "churn_reel": y, "proba_churn": final.predict_proba(X)[:, 1]})
    scores["perte_attendue_fcfa"] = scores["proba_churn"] * scores["MONTHLY_CHARGE"]
    scores["decile_risque"] = pd.qcut(scores["proba_churn"], 10, labels=False,
                                      duplicates="drop") + 1
    return scores, auc
